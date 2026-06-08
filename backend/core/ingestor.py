from __future__ import annotations

import csv
import logging
import os
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pdfplumber
import pytesseract
from PIL import Image


logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}


def _detect_content_type(file_path: str, provided: Optional[str]) -> str:
    if provided:
        return provided
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        return "application/pdf"
    if suffix == ".docx":
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if suffix == ".xlsx":
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if suffix == ".csv":
        return "text/csv"
    if suffix in TEXT_EXTENSIONS:
        return "text/markdown" if suffix in {".md", ".markdown"} else "text/plain"
    if suffix in IMAGE_EXTENSIONS:
        return f"image/{suffix.lstrip('.')}"
    return "application/octet-stream"


def _sliding_chunks(text: str, chunk_size: int, chunk_overlap: int) -> List[Tuple[str, int]]:
    cleaned = re.sub(r"\s+", " ", text).strip()
    if not cleaned:
        return []
    if chunk_size <= 0:
        chunk_size = 1000
    if chunk_overlap >= chunk_size:
        chunk_overlap = max(chunk_size // 4, 1)

    chunks: List[Tuple[str, int]] = []
    start = 0
    while start < len(cleaned):
        end = min(len(cleaned), start + chunk_size)
        chunk = cleaned[start:end].strip()
        if chunk:
            chunks.append((chunk, start))
        if end >= len(cleaned):
            break
        start = max(end - chunk_overlap, start + 1)
    return chunks


def _read_text_file(file_path: str) -> List[Tuple[str, Dict[str, Any]]]:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
        lines = handle.readlines()

    blocks: List[Tuple[str, Dict[str, Any]]] = []
    for index, line in enumerate(lines, start=1):
        stripped = line.strip()
        if stripped:
            blocks.append((stripped, {"line_number": index}))
    return blocks


def _read_pdf_file(file_path: str) -> List[Tuple[str, Dict[str, Any]]]:
    blocks: List[Tuple[str, Dict[str, Any]]] = []
    with pdfplumber.open(file_path) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if not text.strip():
                try:
                    image = page.to_image(resolution=150).original
                    text = pytesseract.image_to_string(image)
                except Exception as exc:
                    logger.warning("OCR fallback failed for PDF page %s in %s: %s", index, file_path, exc)
            if text.strip():
                blocks.append((text.strip(), {"page_number": index}))
    return blocks


def _read_docx_file(file_path: str) -> List[Tuple[str, Dict[str, Any]]]:
    blocks: List[Tuple[str, Dict[str, Any]]] = []
    with zipfile.ZipFile(file_path) as archive:
        xml_bytes = archive.read("word/document.xml")
    root = ET.fromstring(xml_bytes)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    for index, paragraph in enumerate(root.findall(".//w:p", ns), start=1):
        text_parts = [node.text for node in paragraph.findall(".//w:t", ns) if node.text]
        text = "".join(text_parts).strip()
        if text:
            blocks.append((text, {"line_number": index}))
    return blocks


def _read_image_file(file_path: str) -> List[Tuple[str, Dict[str, Any]]]:
    image = Image.open(file_path)
    text = pytesseract.image_to_string(image)
    return [(text.strip(), {"page_number": 1})] if text.strip() else []


def _read_excel_file(file_path: str) -> List[Tuple[str, Dict[str, Any]]]:
    import openpyxl
    blocks: List[Tuple[str, Dict[str, Any]]] = []
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                row_vals = [str(val).strip() for val in row if val is not None]
                if row_vals:
                    text = f"Sheet: {sheet_name} | Row {row_idx}: " + " | ".join(row_vals)
                    blocks.append((text, {"sheet_name": sheet_name, "line_number": row_idx}))
    except Exception as exc:
        logger.exception("Failed to parse Excel file %s: %s", file_path, exc)
    return blocks
def _read_csv_file(file_path: str) -> List[Tuple[str, Dict[str, Any]]]:
    blocks: List[Tuple[str, Dict[str, Any]]] = []
    with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
        reader = csv.reader(handle)
        for index, row in enumerate(reader, start=1):
            text = " ".join(row).strip()
            if text:
                blocks.append((text, {"line_number": index}))
    return blocks


def _load_blocks(file_path: str, content_type: str) -> List[Tuple[str, Dict[str, Any]]]:
    lower_path = file_path.lower()
    if content_type == "application/pdf" or lower_path.endswith(".pdf"):
        return _read_pdf_file(file_path)
    if content_type.endswith("wordprocessingml.document") or lower_path.endswith(".docx"):
        return _read_docx_file(file_path)
    if content_type.endswith("spreadsheetml.sheet") or lower_path.endswith(".xlsx"):
        return _read_excel_file(file_path)
    if content_type == "text/csv" or lower_path.endswith(".csv"):
        return _read_csv_file(file_path)
    if content_type.startswith("image/") or Path(file_path).suffix.lower() in IMAGE_EXTENSIONS:
        return _read_image_file(file_path)
    return _read_text_file(file_path)


def ingest_document(
    file_path: str,
    content_type: Optional[str] = None,
    chunk_size: int = 1000,
    chunk_overlap: int = 200,
) -> List[Dict[str, Any]]:
    """
    Parse a document into evidence-preserving text chunks.
    """
    if not os.path.exists(file_path):
        logger.error("File not found: %s", file_path)
        return []

    resolved_content_type = _detect_content_type(file_path, content_type)

    try:
        blocks = _load_blocks(file_path, resolved_content_type)
    except Exception as exc:
        logger.exception("Failed to ingest %s: %s", file_path, exc)
        return []

    if not blocks:
        return []

    output_chunks: List[Dict[str, Any]] = []
    chunk_index = 0
    for block_text, block_metadata in blocks:
        for chunk_text, _ in _sliding_chunks(block_text, chunk_size=chunk_size, chunk_overlap=chunk_overlap):
            metadata = {
                "source_document": file_path,
                "content_type": resolved_content_type,
                "chunk_index": chunk_index,
                **block_metadata,
            }
            if "page_number" not in metadata and "line_number" not in metadata:
                metadata["line_number"] = 1

            output_chunks.append(
                {
                    "content": chunk_text,
                    "metadata": metadata,
                    "confidence": 1.0,
                }
            )
            chunk_index += 1

    logger.info("Ingested %s into %s chunks", file_path, len(output_chunks))
    return output_chunks
